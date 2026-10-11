#!/usr/bin/env python3
"""Versioned local data library: import XLSX/PDF, attach analysis, export a portable bundle."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, time, timezone
import hashlib
import io
import json
import math
from pathlib import Path
import re
import sqlite3
import sys
import uuid

IMPORTER_VERSION = "1.0.0"
SCHEMA_VERSION = 1
SCHEMA_PATH = Path(__file__).with_name("library_schema.sql")
NAMESPACE = uuid.UUID("da7dd6fc-b6dc-4c46-8fc9-f0e2f329cd65")
SENTIMENTS = {"positive", "negative", "mixed", "neutral", "undetermined"}
TABLES = ("datasets", "dataset_versions", "sheets", "source_rows", "reviews", "review_rows",
          "sales_observations", "study_observations", "document_pages", "import_issues",
          "analysis_runs", "analysis_topics", "topic_reviews", "fragments")


def dumps(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":"))


def stable_id(*parts):
    return str(uuid.uuid5(NAMESPACE, dumps(parts)))


def now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def sha(data):
    return hashlib.sha256(data).hexdigest()


def cell_value(value):
    if isinstance(value, (datetime, date, time)):
        return {"type": type(value).__name__, "iso8601": value.isoformat()}
    return value


def connect(path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(path, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    version = conn.execute("PRAGMA user_version").fetchone()[0]
    if version not in (0, SCHEMA_VERSION):
        conn.close()
        raise ValueError(f"Unsupported database schema {version}; expected {SCHEMA_VERSION}.")
    if version == 0:
        if conn.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchone():
            conn.close()
            raise ValueError("Refusing to initialize an unrelated SQLite database.")
        conn.executescript(SCHEMA_PATH.read_text(encoding="utf-8"))
    return conn


def put(conn, table, **record):
    # Only internal table and column names reach this helper. Values are bound parameters.
    columns = ",".join(record)
    conn.execute(f"INSERT INTO {table} ({columns}) VALUES ({','.join('?' for _ in record)})", tuple(record.values()))


def add_asset(conn, raw, extension):
    digest = sha(raw)
    conn.execute("INSERT OR IGNORE INTO assets VALUES (?,?,?,?)", (digest, extension.lower(), len(raw), raw))
    return digest


def issue(conn, version, code, message, location=None):
    put(conn, "import_issues", issue_id=stable_id(version, code, location, message), version_id=version,
        severity="warning", code=code, location=location, message=message)


def numeric(value, location):
    if type(value) not in (int, float) or not math.isfinite(value):
        raise ValueError(f"{location}: expected a finite numeric value, found {value!r}.")
    return value


def detect_kind(workbook):
    headers = [list(next(w.iter_rows(min_row=1, max_row=1, values_only=True), ())) for w in workbook]
    if any("ID" in h and "Customer Feedback" in h for h in headers):
        return "feedback"
    if headers and all("market_code" in h and any(str(c).startswith("volume_") for c in h) for h in headers):
        return "sales"
    if any(re.fullmatch(r"(?:CN_EU|US)_\d{4}", s) for s in workbook.sheetnames):
        return "studies"
    return "generic"


def save_workbook_rows(conn, version, workbook, cached):
    formula_count = 0
    for sheet in workbook:
        cached_sheet = cached[sheet.title]
        put(conn, "sheets", version_id=version, sheet_name=sheet.title,
            row_count=sheet.max_row, column_count=sheet.max_column,
            metadata_json=dumps({"state": sheet.sheet_state, "merged_ranges": [str(r) for r in sheet.merged_cells]}))
        for row in sheet:
            if all(c.value is None for c in row):
                continue
            values, metadata = [], {}
            for c in row:
                value = c.value
                item = {"data_type": c.data_type, "number_format": c.number_format}
                if c.data_type == "f":
                    formula_count += 1
                    item["formula"] = value
                    value = cached_sheet[c.coordinate].value
                    if value is None:
                        issue(conn, version, "formula_without_cache", "Formula has no cached result; it was not recalculated.", f"{sheet.title}!{c.coordinate}")
                values.append(cell_value(value))
                if c.value is not None:
                    metadata[c.column_letter] = item
            put(conn, "source_rows", version_id=version, sheet_name=sheet.title, row_number=row[0].row,
                values_json=dumps(values), cell_metadata_json=dumps(metadata))
    if formula_count:
        issue(conn, version, "formula_cache", "Formula values are Excel caches and may be stale; original formulas and workbook bytes are retained.")


def ingest_feedback(conn, version, dataset, workbook):
    all_reviews = {}
    for sheet in workbook:
        rows = list(sheet.values)
        headers = rows[0] if rows else ()
        if "ID" not in headers or "Customer Feedback" not in headers:
            issue(conn, version, "unmapped_sheet", "Sheet is retained as raw data; feedback headers were not found.", sheet.title)
            continue
        nonempty = [h for h in headers if h is not None]
        if len(nonempty) != len(set(nonempty)):
            raise ValueError(f"{sheet.title}: duplicate column names.")
        for row_number, row in enumerate(rows[1:], 2):
            if all(v is None for v in row):
                continue
            item = {h: cell_value(v) for h, v in zip(headers, row) if h is not None}
            rid, text = item.get("ID"), item.get("Customer Feedback")
            if rid is None or not str(rid).strip() or not isinstance(text, str) or not text.strip():
                raise ValueError(f"{sheet.title}!{row_number}: missing review ID or feedback text.")
            rid = str(rid).strip()
            if rid in all_reviews and all_reviews[rid]["text"] != text:
                raise ValueError(f"Review ID {rid!r} has conflicting texts; nothing from this file was imported.")
            target = all_reviews.setdefault(rid, {"text": text, "metadata": {}, "rows": []})
            target["rows"].append((sheet.title, row_number))
            for k, v in item.items():
                if k not in ("ID", "Customer Feedback") and v is not None:
                    bucket = target["metadata"].setdefault(k, [])
                    if v not in bucket:
                        bucket.append(v)
    if not all_reviews:
        raise ValueError("No feedback records were found.")
    for rid, record in all_reviews.items():
        review_key = stable_id("review", dataset, rid)
        metadata = {k: vs[0] if len(vs) == 1 else vs for k, vs in record["metadata"].items()}
        put(conn, "reviews", version_id=version, review_key=review_key, source_review_id=rid,
            review_text=record["text"], metadata_json=dumps(metadata))
        for sheet, row in record["rows"]:
            put(conn, "review_rows", version_id=version, review_key=review_key, sheet_name=sheet, row_number=row)


def ingest_sales(conn, version, workbook):
    from openpyxl.utils import get_column_letter
    for sheet in workbook:
        rows = list(sheet.values)
        headers = list(rows[0]) if rows else []
        if "market" not in headers or "market_code" not in headers:
            raise ValueError(f"{sheet.title}: missing market or market_code column.")
        years = [(i, int(str(h)[7:])) for i, h in enumerate(headers) if re.fullmatch(r"volume_\d{4}", str(h))]
        if not years:
            raise ValueError(f"{sheet.title}: no volume_YEAR columns.")
        for row_number, row in enumerate(rows[1:], 2):
            if all(v is None for v in row):
                continue
            market, code = row[headers.index("market")], row[headers.index("market_code")]
            if not market or not code:
                raise ValueError(f"{sheet.title}!{row_number}: market identifiers are missing.")
            for col, year in years:
                cell = f"{get_column_letter(col + 1)}{row_number}"
                value = numeric(row[col], f"{sheet.title}!{cell}")
                if value < 0 or int(value) != value:
                    raise ValueError(f"{sheet.title}!{cell}: volume must be a non-negative integer.")
                put(conn, "sales_observations", observation_id=stable_id(version, sheet.title, cell),
                    version_id=version, sheet_name=sheet.title, row_number=row_number, cell_address=cell,
                    model_group=sheet.title, market=str(market), market_code=str(code), year=year, volume=int(value))
    issue(conn, version, "sales_scenario_unspecified", "The source does not label actuals versus forecasts. Scenario remains unspecified, including 2030. Combined model groups are not split.")


def ingest_studies(conn, version, workbook):
    from openpyxl.utils import get_column_letter
    imported = 0
    for sheet in workbook:
        match = re.fullmatch(r"(CN_EU|US)_(\d{4})", sheet.title)
        if not match:
            issue(conn, version, "unmapped_sheet", "Unknown study layout; sheet is retained as raw rows.", sheet.title)
            continue
        region_group, year = match.group(1), int(match.group(2))
        rows = list(sheet.values)
        header_index = 2 if region_group == "CN_EU" else 0
        if len(rows) <= header_index:
            raise ValueError(f"{sheet.title}: study header is missing.")
        columns = []
        region = None
        for col, value in enumerate(rows[header_index]):
            if region_group == "CN_EU" and rows[1][col]:
                region = {"China": "CN", "EU": "EU"}.get(str(rows[1][col]).strip())
            model = re.search(r"\b([FG]\d{2})\b", str(value or ""))
            if model:
                if region_group == "CN_EU" and region is None:
                    raise ValueError(f"{sheet.title}: cannot determine region for {value}.")
                columns.append((col, model.group(1), str(value).strip(), region if region_group == "CN_EU" else "US"))
        if not columns:
            raise ValueError(f"{sheet.title}: model columns were not found.")
        metric = str(rows[header_index][0]).strip()
        for row_number, row in enumerate(rows[header_index + 1:], header_index + 2):
            if all(v is None for v in row):
                continue
            label = row[0]
            if not any(row[c] is not None for c, *_ in columns):
                if label:
                    metric = str(label).strip()
                continue
            if not label or not metric:
                raise ValueError(f"{sheet.title}!{row_number}: metric or response label is missing.")
            if region_group == "CN_EU" and label != "Mean":
                raise ValueError(f"{sheet.title}!{row_number}: expected a Mean row.")
            if region_group == "US" and row[1] != "% of column, sample base":
                raise ValueError(f"{sheet.title}!{row_number}: unsupported study statistic.")
            for col, model, model_label, region in columns:
                cell = f"{get_column_letter(col + 1)}{row_number}"
                if row[col] is None:
                    issue(conn, version, "missing_study_value", "Missing study observation; no zero was substituted.", f"{sheet.title}!{cell}")
                    continue
                value = numeric(row[col], f"{sheet.title}!{cell}")
                if region_group == "US" and not 0 <= value <= 1:
                    raise ValueError(f"{sheet.title}!{cell}: expected a fraction from zero to one.")
                put(conn, "study_observations", observation_id=stable_id(version, sheet.title, cell),
                    version_id=version, sheet_name=sheet.title, row_number=row_number, cell_address=cell,
                    metric=metric, response_label=str(label), model_code=model, model_label=model_label,
                    region_code=region, year=year, statistic="mean" if region_group == "CN_EU" else "column_share",
                    value=value, unit="unspecified_scale" if region_group == "CN_EU" else "fraction", sample_size=None)
                imported += 1
    if not imported:
        raise ValueError("No recognized study observations were found; use --kind generic to archive another layout.")
    issue(conn, version, "study_scale_and_sample_size", "Mean scale bounds and respondent counts are not specified. A Sample total value of 1 is a column share, not one respondent. Do not average across regions or studies without comparable measures and sample sizes.")


def ingest_file(conn, path, dataset_key=None, kind="auto"):
    path = Path(path)
    if path.suffix.lower() not in (".xlsx", ".pdf"):
        raise ValueError("Supported source formats are .xlsx and .pdf.")
    raw = path.read_bytes()
    key = dataset_key or path.stem.casefold()
    if not key.strip():
        raise ValueError("Dataset key must not be empty.")
    dataset = stable_id("dataset", key)
    version = stable_id("version", dataset, sha(raw), IMPORTER_VERSION)
    previous = conn.execute("SELECT version_id FROM dataset_versions WHERE version_id=?", (version,)).fetchone()
    if previous:
        existing_kind = conn.execute("SELECT kind FROM datasets WHERE dataset_id=?", (dataset,)).fetchone()[0]
        if kind != "auto" and kind != existing_kind:
            raise ValueError("This dataset already exists with a different kind.")
        return {"status": "unchanged", "dataset_key": key, "version_id": version}
    workbook = cached = None
    try:
        if path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(raw), data_only=False)
            cached = load_workbook(io.BytesIO(raw), data_only=True)
            chosen_kind = detect_kind(cached) if kind == "auto" else kind
            if chosen_kind == "document":
                raise ValueError("Document imports require a PDF.")
        else:
            chosen_kind = "document"
            if kind not in ("auto", "document"):
                raise ValueError("PDF files must use kind document.")
        with conn:
            existing = conn.execute("SELECT kind FROM datasets WHERE dataset_id=?", (dataset,)).fetchone()
            if existing and existing[0] != chosen_kind:
                raise ValueError("Dataset kind changed. Use a different --dataset key or the previous layout.")
            if not existing:
                put(conn, "datasets", dataset_id=dataset, dataset_key=key, kind=chosen_kind, created_at=now())
            number = conn.execute("SELECT COALESCE(MAX(version_number),0)+1 FROM dataset_versions WHERE dataset_id=?", (dataset,)).fetchone()[0]
            digest = add_asset(conn, raw, path.suffix)
            put(conn, "dataset_versions", version_id=version, dataset_id=dataset, version_number=number,
                asset_sha256=digest, original_filename=path.name, imported_at=now(), importer_version=IMPORTER_VERSION,
                metadata_json=dumps({"kind": chosen_kind, "source_extension": path.suffix.lower()}))
            if workbook is not None:
                save_workbook_rows(conn, version, workbook, cached)
                if chosen_kind == "feedback":
                    ingest_feedback(conn, version, dataset, cached)
                elif chosen_kind == "sales":
                    ingest_sales(conn, version, cached)
                elif chosen_kind == "studies":
                    ingest_studies(conn, version, cached)
                elif chosen_kind == "generic":
                    issue(conn, version, "generic_layout", "Only raw rows are available. Add a dedicated adapter before treating these rows as normalized observations.")
                else:
                    raise ValueError(f"Unsupported dataset kind: {chosen_kind}")
            else:
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(raw))
                if reader.is_encrypted and not reader.decrypt(""):
                    raise ValueError("Password-protected PDFs require an accessible source copy.")
                for page_number, page in enumerate(reader.pages, 1):
                    text = page.extract_text() or ""
                    put(conn, "document_pages", version_id=version, page_number=page_number,
                        extracted_text=text, extraction_status="text_extracted" if text.strip() else "needs_ocr")
                    if not text.strip():
                        issue(conn, version, "needs_ocr", "This page has no extractable text; original page is retained in the PDF.", f"page {page_number}")
                issue(conn, version, "pdf_layout", "PDF text is indexed by page. Table layout, prices, option codes and availability relationships have not been normalized or verified. Consult the original page.")
            counts = {t: conn.execute(f"SELECT COUNT(*) FROM {t} WHERE version_id=?", (version,)).fetchone()[0]
                      for t in ("source_rows", "reviews", "sales_observations", "study_observations", "document_pages", "import_issues")}
        return {"status": "imported", "dataset_key": key, "version": number, "kind": chosen_kind, **counts}
    finally:
        if workbook:
            workbook.close()
        if cached:
            cached.close()


def attach_analysis(conn, topics_path, summary_path, dataset_key=None):
    import summarize_topics
    from export_topics import validate
    topics_raw, summary_raw = Path(topics_path).read_bytes(), Path(summary_path).read_bytes()
    source, summary = json.loads(topics_raw.decode("utf-8-sig")), json.loads(summary_raw.decode("utf-8-sig"))
    validate(summary)
    summarize_topics.validate_source(source)
    if summary["source"]["input_sha256"] != sha(topics_raw):
        raise ValueError("Summary does not match the supplied stage-1 JSON checksum.")
    candidates = conn.execute("SELECT v.*,d.dataset_key,d.kind FROM dataset_versions v JOIN datasets d USING(dataset_id) WHERE v.asset_sha256=?",
                              (source["source"]["sha256"],)).fetchall()
    candidates = [c for c in candidates if c["kind"] == "feedback" and (not dataset_key or c["dataset_key"] == dataset_key)]
    if len(candidates) != 1:
        raise ValueError("Analysis must match exactly one imported feedback version; import its workbook or specify --dataset.")
    version = candidates[0]["version_id"]
    reviews = {r["source_review_id"]: r for r in conn.execute("SELECT * FROM reviews WHERE version_id=?", (version,))}
    source_reviews = {r["id"]: r for r in source["reviews"]}
    if set(source_reviews) != set(reviews) or any(r["text"] != reviews[rid]["review_text"] for rid, r in source_reviews.items()):
        raise ValueError("Analysis review identities or texts differ from the imported workbook.")
    if set(summary["topics"]) != {t["topic_id"] for t in source["topics"]}:
        raise ValueError("Stage-1 and summary topic IDs differ.")
    if summary["overview"]["review_count"] != len(reviews):
        raise ValueError("Summary overview count differs from source reviews.")
    run = stable_id("analysis", version, sha(topics_raw), sha(summary_raw))
    if conn.execute("SELECT 1 FROM analysis_runs WHERE run_id=?", (run,)).fetchone():
        return {"status": "unchanged", "analysis_run_id": run}
    with conn:
        put(conn, "analysis_runs", run_id=run, version_id=version, topics_asset_sha256=add_asset(conn, topics_raw, ".json"),
            summary_asset_sha256=add_asset(conn, summary_raw, ".json"), attached_at=now(), summary_status=summary["status"],
            metadata_json=dumps({k: v for k, v in summary.items() if k not in ("topics", "review_metadata")}))
        for t in source["topics"]:
            tid = t["topic_id"]
            entry = summary["topics"][tid]
            assignment_ids = {a["review_id"] for a in t["reviews"]}
            counts = Counter(a["sentiment"] for a in t["reviews"])
            if (assignment_ids != set(entry["review_ids"]) or len(assignment_ids) != entry["review_count"] or
                    any(counts[s] != entry["sentiment"][s]["count"] for s in SENTIMENTS)):
                raise ValueError(f"{tid}: summary counts or assignments do not match stage 1.")
            put(conn, "analysis_topics", run_id=run, topic_id=tid, title=entry["title"],
                review_count=entry["review_count"], data_json=dumps(entry))
            for a in t["reviews"]:
                put(conn, "topic_reviews", run_id=run, topic_id=tid, version_id=version,
                    review_key=reviews[a["review_id"]]["review_key"], sentiment=a["sentiment"])
        for f in source["fragments"]:
            put(conn, "fragments", run_id=run, fragment_id=f["fragment_id"], version_id=version,
                review_key=reviews[f["review_id"]]["review_key"], data_json=dumps(f))
    return {"status": "attached", "analysis_run_id": run, "topics": len(source["topics"]), "fragments": len(source["fragments"])}


def inventory(conn):
    records = [dict(r) for r in conn.execute("""SELECT d.dataset_key,d.kind,v.version_number,v.original_filename,v.version_id,
      (SELECT COUNT(*) FROM reviews r WHERE r.version_id=v.version_id) AS reviews,
      (SELECT COUNT(*) FROM sales_observations r WHERE r.version_id=v.version_id) AS sales,
      (SELECT COUNT(*) FROM study_observations r WHERE r.version_id=v.version_id) AS studies,
      (SELECT COUNT(*) FROM document_pages r WHERE r.version_id=v.version_id) AS pages
      FROM datasets d JOIN current_versions v USING(dataset_id) ORDER BY d.dataset_key""")]
    for record in records:
        run = conn.execute("""SELECT run_id,summary_status FROM analysis_runs WHERE version_id=?
            ORDER BY (summary_status='complete') DESC, attached_at DESC, rowid DESC LIMIT 1""",
            (record["version_id"],)).fetchone()
        record["analysis_status"] = ("not_applicable" if record["kind"] != "feedback" else
                                     "awaiting_analysis" if run is None else
                                     "ready" if run["summary_status"] == "complete" else run["summary_status"])
        record["analysis_run_id"] = run["run_id"] if run else None
        record["analysis_topics"] = conn.execute("SELECT COUNT(*) FROM analysis_topics WHERE run_id=?",
                                                (run["run_id"],)).fetchone()[0] if run else 0
    return records


def status_markdown(conn):
    from export_topics import md
    records = inventory(conn)
    lines = ["# Data library status", "", f"Generated: {now()}", "",
             "Only the current source version is shown. A completed analysis is preferred over partial runs.",
             "Ready means processing is complete, not that AI conclusions have been independently validated.", ""]
    for heading, selected in (
        ("Analysis ready", [r for r in records if r["analysis_status"] == "ready"]),
        ("Awaiting or incomplete analysis", [r for r in records if r["kind"] == "feedback" and r["analysis_status"] != "ready"]),
        ("Supporting data imported", [r for r in records if r["kind"] != "feedback"]),
    ):
        lines += [f"## {heading}", "", "| Dataset | Version | Reviews | Topics | Status |",
                  "| --- | --- | --- | --- | --- |"]
        for r in selected:
            status = (f"Imported: {r['sales']} sales, {r['studies']} study observations, {r['pages']} PDF pages"
                      if r["kind"] != "feedback" else r["analysis_status"].replace("_", " ").capitalize())
            lines.append(f"| {md(r['original_filename'])} | {r['version_number']} | {r['reviews']} | {r['analysis_topics']} | {md(status)} |")
        if not selected:
            lines.append("| None | | | | |")
        lines.append("")
    lines += ["## Refresh after processing another workbook", "", "Run from the feedback_topics_package directory:",
              "", "```powershell", "python data_library.py ingest-all .", "python data_library.py status", "```", "",
              "This attaches matching existing analysis files; it does not run AI analysis.",
              "Status files and export bundles are snapshots. Regenerate them after updates.", ""]
    return "\n".join(lines)


def check(conn):
    integrity = conn.execute("PRAGMA integrity_check").fetchall()
    if [r[0] for r in integrity] != ["ok"] or conn.execute("PRAGMA foreign_key_check").fetchone():
        raise ValueError("Database integrity or reference check failed.")
    for asset in conn.execute("SELECT * FROM assets"):
        if len(asset["content"]) != asset["byte_size"] or sha(asset["content"]) != asset["sha256"]:
            raise ValueError("Stored source checksum mismatch.")
    return {"status": "ok", "datasets": len(inventory(conn)),
            "counts": {t: conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in TABLES}}


def export_bundle(conn, output):
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    if any(output.iterdir()):
        raise ValueError("Export directory must be empty; snapshots never overwrite earlier exports.")
    # Work from a consistent SQLite backup rather than copying an active database file.
    snapshot_path = output / "library.sqlite"
    snapshot = sqlite3.connect(snapshot_path)
    try:
        conn.backup(snapshot)
        snapshot.row_factory = sqlite3.Row
        validation = check(snapshot)
        files = []

        def record(path, count=None):
            raw = path.read_bytes()
            files.append({"path": path.relative_to(output).as_posix(), "sha256": sha(raw), "bytes": len(raw), "records": count})

        for table in TABLES:
            # Export JSON columns as objects, retaining their explicit _json column names.
            path = output / f"{table}.jsonl"
            count = 0
            with path.open("w", encoding="utf-8", newline="\n") as stream:
                for row in snapshot.execute(f"SELECT * FROM {table} ORDER BY rowid"):
                    item = {k: json.loads(row[k]) if k.endswith("_json") else row[k] for k in row.keys()}
                    stream.write(dumps(item) + "\n")
                    count += 1
            record(path, count)
        (output / "assets").mkdir()
        asset_records = []
        for a in snapshot.execute("SELECT * FROM assets ORDER BY sha256"):
            path = output / "assets" / f"{a['sha256']}{a['extension']}"
            path.write_bytes(a["content"])
            record(path)
            asset_records.append({"sha256": a["sha256"], "byte_size": a["byte_size"], "extension": a["extension"], "path": path.relative_to(output).as_posix()})
        asset_index = output / "assets.jsonl"
        asset_index.write_text("".join(dumps(a) + "\n" for a in asset_records), encoding="utf-8")
        record(asset_index, len(asset_records))
        schema = output / "schema.sql"
        schema.write_text(SCHEMA_PATH.read_text(encoding="utf-8"), encoding="utf-8")
        record(schema)
        dictionary = {"schema_version": SCHEMA_VERSION, "tables": {}}
        for table in ("assets", *TABLES):
            dictionary["tables"][table] = {"columns": [dict(r) for r in snapshot.execute(f"PRAGMA table_info({table})")],
                                           "foreign_keys": [dict(r) for r in snapshot.execute(f"PRAGMA foreign_key_list({table})")]}
        path = output / "schema.json"
        path.write_text(json.dumps(dictionary, indent=2), encoding="utf-8")
        record(path)
        guide = output / "README.md"
        guide.write_text(Path(__file__).with_name("DATA_LIBRARY.md").read_text(encoding="utf-8"), encoding="utf-8")
        record(guide)
        catalog = output / "STATUS.md"
        catalog.write_text(status_markdown(snapshot), encoding="utf-8")
        record(catalog)
        snapshot.close()
        record(snapshot_path)
        manifest = {"format": "kolosok-data-library", "schema_version": SCHEMA_VERSION, "created_at": now(),
                    "version_scope": "all_versions", "current_version_rule": "MAX(version_number) per dataset_id",
                    "jsonl_json_columns": "Columns ending in _json are nested objects in JSONL; serialize them as JSON text when importing into SQL.",
                    "validation": validation, "files": files}
        (output / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
        return {"status": "exported", "output": str(output), "files": len(files) + 1}
    finally:
        snapshot.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=Path("data_library/library.sqlite"))
    sub = parser.add_subparsers(dest="command", required=True)
    ingest = sub.add_parser("ingest", help="Import one source file; identical content is skipped")
    ingest.add_argument("input", type=Path)
    ingest.add_argument("--dataset", help="Stable logical key; default is the lower-case filename stem")
    ingest.add_argument("--kind", choices=("auto", "feedback", "sales", "studies", "document", "generic"), default="auto")
    all_files = sub.add_parser("ingest-all", help="Import all XLSX/PDF files in a folder, then attach matching analyses")
    all_files.add_argument("folder", type=Path)
    attach = sub.add_parser("attach-analysis")
    attach.add_argument("--topics", type=Path, required=True)
    attach.add_argument("--summary", type=Path, required=True)
    attach.add_argument("--dataset")
    sub.add_parser("list")
    sub.add_parser("check")
    status = sub.add_parser("status", help="Write a readable status file with completed analyses listed separately")
    status.add_argument("--output", type=Path, help="Default: STATUS.md beside the database")
    export = sub.add_parser("export")
    export.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    conn = None
    try:
        conn = connect(args.db)
        if args.command == "ingest":
            result = ingest_file(conn, args.input, args.dataset, args.kind)
        elif args.command == "ingest-all":
            if not args.folder.is_dir():
                raise ValueError("Input folder does not exist.")
            results, failures = [], []
            for p in sorted(args.folder.iterdir()):
                if p.is_file() and not p.name.startswith("~$") and p.suffix.lower() in (".xlsx", ".pdf"):
                    try:
                        results.append(ingest_file(conn, p))
                    except Exception as error:
                        failures.append({"file": p.name, "error": str(error)})
            for p in sorted(args.folder.glob("*_topics.json")):
                summary = p.with_name(p.stem + "_summary.json")
                if summary.is_file():
                    try:
                        results.append(attach_analysis(conn, p, summary))
                    except Exception as error:
                        failures.append({"file": p.name, "error": str(error)})
            print(json.dumps({"results": results, "failures": failures}, ensure_ascii=True, indent=2))
            return 2 if failures else 0
        elif args.command == "attach-analysis":
            result = attach_analysis(conn, args.topics, args.summary, args.dataset)
        elif args.command == "list":
            result = inventory(conn)
        elif args.command == "check":
            result = check(conn)
        elif args.command == "status":
            output = args.output or args.db.with_name("STATUS.md")
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text(status_markdown(conn), encoding="utf-8")
            result = {"status": "written", "output": str(output.resolve())}
        else:
            result = export_bundle(conn, args.output)
        print(json.dumps(result, ensure_ascii=True, indent=2))
        return 0
    except (ValueError, OSError, sqlite3.Error, ImportError, KeyError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 2
    finally:
        if conn is not None:
            conn.close()


if __name__ == "__main__":
    raise SystemExit(main())
