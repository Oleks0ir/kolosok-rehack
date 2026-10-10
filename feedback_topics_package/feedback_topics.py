#!/usr/bin/env python3
"""Local, deterministic TF-IDF + average-linkage clustering of Excel feedback.

No model downloads, neural networks, external APIs, scipy or scikit-learn.
Dependencies: numpy, openpyxl. See README_RU.md for definitions and limitations.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
import re
import sys
import time
from zipfile import BadZipFile

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

VERSION = "1.0.0"
STOPWORDS = set("""
a an the this that these those i me my mine we us our ours you your yours he him
his she her hers it its they them their theirs myself yourself himself herself
itself ourselves themselves am is are was were be been being do does did doing
have has had having can could may might must shall should will would to of in
on at by for from with as into through about over under up down out off than
then so such very really quite much more most less least also just only even
still already now there here when where why how what which who whom whose
and or but if because while although however though despite before after
since until once again ever never not no nor without all any both each every
few many other another own same some too either neither otherwise overall
car cars vehicle vehicles bmw series model models f70 new old previous one
first second get gets got getting go going comes come use used using take
make makes made feel feels think thought know say said lot lots bit little
good great excellent fantastic amazing brilliant nice lovely wonderful bad
poor terrible awful worst best love loves loved like likes liked dislike
dislikes hate hates happy unhappy pleased disappointed recommend definitely
absolutely extremely pretty generally probably personally view experience
im ive id youre youve theyre weve were isnt arent wasnt werent dont doesnt
didnt havent hasnt wouldnt couldnt shouldnt thats theres whats whos bmws
well really actually simply everything anything something nothing thing things
seems seem felt feeling find found want wanted bought buy buying purchased
purchase owned owning years year months month days day week weeks time times
always often sometimes usually last next yet enough lot overall basically
""".split())

# These are word-normalization rules, not a predefined list of output topics.
ALIASES = {
    "seats": "seat", "seating": "seat", "seatback": "backrest",
    "comfortable": "comfort", "uncomfortable": "comfort", "comfy": "comfort",
    "discomfort": "comfort", "comfortably": "comfort", "uncomfortably": "comfort",
    "brakes": "brake", "braking": "brake", "braked": "brake",
    "driving": "drive", "drives": "drive", "driven": "drive",
    "steered": "steering", "screens": "screen", "touchscreens": "touchscreen",
    "displays": "display", "controls": "control", "controlling": "control",
    "buttons": "button", "tyres": "tyre", "tires": "tyre", "tire": "tyre",
    "wheels": "wheel", "windows": "window", "mirrors": "mirror",
    "doors": "door", "lights": "light", "headlights": "headlight",
    "headlamps": "headlight", "headlamp": "headlight",
    "engines": "engine", "gearboxes": "gearbox", "gears": "gear",
    "accelerating": "acceleration", "accelerates": "acceleration",
    "manuals": "manual", "instructions": "instruction",
    "subscriptions": "subscription", "payments": "payment",
    "charges": "charge", "charging": "charge", "charged": "charge",
    "dealers": "dealer", "dealership": "dealer", "dealerships": "dealer",
    "apps": "app", "applications": "app", "application": "app",
    "phones": "phone", "smartphone": "phone", "smartphones": "phone",
    "sensors": "sensor", "warnings": "warning", "alerts": "alert",
    "features": "feature", "functions": "function", "settings": "setting",
    "materials": "material", "plastics": "plastic", "updates": "update",
    "problems": "problem", "issues": "issue", "faults": "fault",
    "noises": "noise", "noisy": "noise", "rattles": "rattle",
    "rattling": "rattle", "squeaks": "squeak", "squeaking": "squeak",
    "creaks": "creak", "creaking": "creak", "vibrations": "vibration",
    "vibrating": "vibration", "heated": "heating", "adjustments": "adjustment",
    "adjusting": "adjustment", "adjustable": "adjustment", "adjust": "adjustment",
    "economical": "economy", "efficient": "efficiency", "inefficient": "efficiency",
    "reliability": "reliability", "reliable": "reliability", "unreliable": "reliability",
    "connects": "connection", "connecting": "connection", "connectivity": "connection",
    "disconnected": "connection", "disconnects": "connection",
    "built": "build", "building": "build", "looks": "look", "looking": "look",
    "enjoyable": "fun", "enjoyment": "fun", "pleasure": "fun", "pleasant": "fun",
}
TOKEN = re.compile(r"[^\W\d_]+(?:['’][^\W\d_]+)?", re.UNICODE)
BOUNDARY = re.compile(
    r"[.!?;\n]+|,(?=\s+(?:and\s+)?(?:the|my|our|its|their|there|it|i)\b)"
    r"|\b(?:but|however|although|though|nevertheless|whereas|on the other hand)\b"
    r"|\band(?=\s+(?:the|my|our|its|their|there)\b)",
    re.IGNORECASE,
)
SENTIMENTS = ("positive", "negative", "mixed", "neutral", "undetermined")


def log(message):
    print(message, file=sys.stderr, flush=True)


def scalar(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def read_reviews(path, sheet_name=None, id_column="ID", text_column="Customer Feedback"):
    """Deduplicate IDs while preserving all metadata values and original row numbers."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = wb[sheet_name] if sheet_name else wb.worksheets[0]
        name = sheet.title
        iterator = sheet.iter_rows(values_only=True)
        first = next(iterator, None)
        if first is None:
            raise ValueError("The selected worksheet is empty.")
        headers = [str(v).strip() if v is not None else "" for v in first]
        for needed in (id_column, text_column):
            if headers.count(needed) != 1:
                raise ValueError(f"Expected exactly one column {needed!r}; found {headers.count(needed)}.")
        named = [h for h in headers if h]
        if len(named) != len(set(named)):
            raise ValueError("Duplicate column names are not supported.")
        reviews, source_rows = {}, 0
        for row_number, row in enumerate(iterator, 2):
            if not any(v is not None for v in row):
                continue
            source_rows += 1
            data = {h: scalar(v) for h, v in zip(headers, row) if h}
            raw_id = data.get(id_column)
            if raw_id is None or not str(raw_id).strip():
                raise ValueError(f"Missing review ID at Excel row {row_number}.")
            rid = str(raw_id).strip()
            text = str(data.get(text_column) or "")
            if rid not in reviews:
                reviews[rid] = {"id": rid, "text": text, "source_rows": [], "metadata": {}}
            review = reviews[rid]
            if review["text"] != text:
                raise ValueError(f"ID {rid!r} has different review texts (row {row_number}). Use a unique ID column.")
            review["source_rows"].append(row_number)
            for key, value in data.items():
                if key in (id_column, text_column):
                    continue
                values = review["metadata"].setdefault(key, [])
                if value not in values:
                    values.append(value)
        # Single values stay scalar. Multiple source labels stay lists, not an arbitrary first value.
        for review in reviews.values():
            review["metadata"] = {k: vs[0] if len(vs) == 1 else vs for k, vs in review["metadata"].items()}
        return list(reviews.values()), name, source_rows
    finally:
        wb.close()


def topic_tokens(text):
    lowered = text.lower().replace("’", "'")
    lowered = re.sub(r"\b([a-z]+)(?:'m|'ve|'re|'ll|'d|'s)\b", r"\1", lowered)
    lowered = re.sub(r"\b([a-z]+)n't\b", r"\1 not", lowered)
    lowered = re.sub(r"\btouch\s+screen\b", "touchscreen", lowered)
    lowered = re.sub(r"\bi\s*drive\b", "idrive", lowered)
    tokens = []
    for match in TOKEN.finditer(lowered):
        word = match.group(0)
        if word in STOPWORDS or len(word) < 2:
            continue
        word = ALIASES.get(word, word)
        if word not in STOPWORDS:
            tokens.append(word)
    return tokens


def split_fragments(reviews, max_words=65):
    """Retain exact substrings/offsets; identical fragments share a clustering observation."""
    nodes, by_normalized_text = [], {}
    occurrences = []
    for review in reviews:
        text = review["text"]
        ends = [(m.start(), m.end()) for m in BOUNDARY.finditer(text)] + [(len(text), len(text))]
        start = 0
        for end, next_start in ends:
            raw = text[start:end]
            # Bound very long clauses without dropping their content.
            words = list(re.finditer(r"\S+", raw))
            for offset in range(0, len(words), max_words):
                selected = words[offset:offset + max_words]
                a, b = start + selected[0].start(), start + selected[-1].end()
                piece = text[a:b]
                if not TOKEN.search(piece):
                    continue
                key = re.sub(r"\s+", " ", piece).strip().lower()
                if key not in by_normalized_text:
                    by_normalized_text[key] = len(nodes)
                    nodes.append({"text": piece, "tokens": topic_tokens(piece), "occurrences": []})
                node_index = by_normalized_text[key]
                occurrence = {
                    "fragment_id": f"F{len(occurrences) + 1:06d}", "review_id": review["id"],
                    "text": piece, "start": a, "end": b, "node_index": node_index,
                }
                nodes[node_index]["occurrences"].append(len(occurrences))
                occurrences.append(occurrence)
            start = next_start
    return nodes, occurrences


def tfidf(nodes, min_df=2, max_df_ratio=0.75, max_features=10000):
    """Sublinear TF, smoothed IDF, L2 normalization; no trained/pretrained models."""
    counts, document_frequency = [], Counter()
    for node in nodes:
        tokens = node["tokens"]
        terms = tokens + [a + " " + b for a, b in zip(tokens, tokens[1:])]
        counter = Counter(terms)
        counts.append(counter)
        document_frequency.update(counter.keys())
    n = len(nodes)
    max_df = math.floor(max_df_ratio * n)
    candidates = [t for t, df in document_frequency.items() if min_df <= df <= max_df]
    candidates.sort(key=lambda t: (-document_frequency[t], t))
    vocabulary = sorted(candidates[:max_features])
    lookup = {term: i for i, term in enumerate(vocabulary)}
    idf = {t: 1.0 + math.log((1.0 + n) / (1.0 + document_frequency[t])) for t in vocabulary}
    vectors = []
    for counter in counts:
        vector = {lookup[t]: (1.0 + math.log(c)) * idf[t] for t, c in counter.items() if t in lookup}
        norm = math.sqrt(sum(v * v for v in vector.values()))
        vectors.append({k: v / norm for k, v in vector.items()} if norm else {})
    return vectors, vocabulary


def cosine_matrix(vectors, max_memory_mb=768):
    """Exact sparse dot products accumulated into a dense float64 similarity matrix."""
    n = len(vectors)
    matrix_bytes = n * n * 8
    # Include temporary outer products and advanced-indexing copies conservatively.
    estimate = matrix_bytes * 1.3 + n * 8 * 256 * 3
    if estimate > max_memory_mb * 1024**2:
        raise ValueError(
            f"{n} nonempty unique fragments require approximately {estimate / 1024**2:.0f} MB "
            f"for similarity computation, exceeding --max-memory-mb={max_memory_mb}. "
            "Increase that limit if RAM permits. No reviews were silently sampled or dropped."
        )
    similarities = np.zeros((n, n), dtype=np.float64)
    postings = defaultdict(list)
    for i, vector in enumerate(vectors):
        for term, value in vector.items():
            postings[term].append((i, value))
    for entries in postings.values():
        ids = np.fromiter((i for i, _ in entries), dtype=np.intp)
        weights = np.fromiter((v for _, v in entries), dtype=np.float64)
        # Chunk the temporary term-by-term outer product, preserving exact dot products.
        for start in range(0, len(ids), 256):
            block = ids[start:start + 256]
            ix = np.ix_(block, ids)
            similarities[ix] += weights[start:start + 256, None] * weights[None, :]
    np.clip(similarities, 0.0, 1.0, out=similarities)
    np.fill_diagonal(similarities, -np.inf)
    return similarities


def average_linkage(similarities, distance_threshold):
    """Nearest-neighbor-chain average linkage, cut at the requested cosine distance.

    The input symmetric similarity matrix is consumed. For reciprocal nearest
    neighbors A,B, S(A union B,C)=(|A|S(A,C)+|B|S(B,C))/(|A|+|B|).
    A cluster whose best similarity is below the cut can be finalized: averages
    of its remaining similarities can never exceed their current maximum.
    This is an O(n^2)-memory, O(n^2)-time reducible-linkage algorithm.
    """
    n = len(similarities)
    if not n:
        return []
    cutoff = 1.0 - distance_threshold
    active = np.ones(n, dtype=bool)
    sizes = np.ones(n, dtype=np.int64)
    members = {i: [i] for i in range(n)}
    chain, result = [], []
    remaining = n
    while remaining:
        if not chain:
            chain.append(int(np.flatnonzero(active)[0]))
        current = chain[-1]
        neighbor = int(np.argmax(similarities[current]))
        best = similarities[current, neighbor]
        if not np.isfinite(best) or best < cutoff - 1e-12:
            result.append(sorted(members.pop(current)))
            active[current] = False
            similarities[current, :] = -np.inf
            similarities[:, current] = -np.inf
            chain.pop()
            remaining -= 1
            continue
        # Prefer the predecessor on ties; this prevents longer equal-distance cycles.
        if len(chain) > 1 and similarities[current, chain[-2]] >= best - 1e-14:
            neighbor = chain[-2]
        if len(chain) > 1 and neighbor == chain[-2]:
            keep, drop = min(current, neighbor), max(current, neighbor)
            combined = (similarities[keep] * sizes[keep] + similarities[drop] * sizes[drop]) / (sizes[keep] + sizes[drop])
            similarities[keep, :] = combined
            similarities[:, keep] = combined
            similarities[drop, :] = -np.inf
            similarities[:, drop] = -np.inf
            similarities[keep, keep] = -np.inf
            members[keep].extend(members.pop(drop))
            sizes[keep] += sizes[drop]
            active[drop] = False
            remaining -= 1
            del chain[-2:]
        else:
            chain.append(neighbor)
    return sorted(result, key=lambda group: group[0])


def normalize_sentiment_text(text):
    text = text.lower().replace("’", "'")
    replacements = {
        "can't": "can not", "cannot": "can not", "won't": "will not",
        "don't": "do not", "dont": "do not", "doesn't": "does not", "doesnt": "does not",
        "didn't": "did not", "didnt": "did not", "isn't": "is not", "isnt": "is not",
        "aren't": "are not", "arent": "are not", "wasn't": "was not", "wasnt": "was not",
        "weren't": "were not", "hasn't": "has not", "haven't": "have not",
        "wouldn't": "would not", "couldn't": "could not", "shouldn't": "should not",
    }
    for old, new in replacements.items():
        text = re.sub(r"(?<!\w)" + re.escape(old) + r"(?!\w)", new, text)
    return text


class SentimentRules:
    def __init__(self, configuration):
        self.configuration = configuration
        self.rules = []
        for rule in configuration["rules"]:
            self.rules.append((rule, re.compile(rule["pattern"], re.I)))
        self.request = re.compile(configuration["request_pattern"], re.I)
        self.negations = set(configuration["negations"])

    def analyze(self, text):
        normalized = normalize_sentiment_text(text)
        matches = []
        for rule, pattern in self.rules:
            for match in pattern.finditer(normalized):
                matches.append((match.start(), match.end(), rule, match.group(0)))
        # Longer expressions take precedence over overlapping individual words.
        matches.sort(key=lambda x: (-(x[1] - x[0]), x[0], x[2]["id"]))
        selected = []
        for a, b, rule, expression in matches:
            if any(a < y and b > x for x, y, *_ in selected):
                continue
            selected.append((a, b, rule, expression))
        evidence = []
        for a, b, rule, expression in sorted(selected, key=lambda x: x[0]):
            score = float(rule["score"])
            negated = False
            if rule.get("allow_negation", True):
                prefix = re.split(r"[.!?;,]", normalized[:a])[-1]
                preceding = [m.group(0) for m in TOKEN.finditer(prefix)][-self.configuration["negation_window"]:]
                # "not only comfortable" is not a negation of comfortable.
                preceding = [w for i, w in enumerate(preceding) if not (w == "not" and i + 1 < len(preceding) and preceding[i + 1] == "only")]
                negated = sum(w in self.negations for w in preceding) % 2 == 1
                if negated:
                    score *= -1
            evidence.append({"rule_id": rule["id"], "expression": expression, "score": score, "negated": negated})
        pos = sum(e["score"] for e in evidence if e["score"] > 0)
        neg = -sum(e["score"] for e in evidence if e["score"] < 0)
        label = "mixed" if pos and neg else "positive" if pos else "negative" if neg else "undetermined"
        return {"label": label, "positive_score": pos, "negative_score": neg,
                "request": bool(self.request.search(normalized)), "evidence": evidence}


def combined_sentiment(labels):
    labels = set(labels)
    if "mixed" in labels or {"positive", "negative"} <= labels:
        return "mixed"
    for label in ("positive", "negative", "neutral"):
        if label in labels:
            return label
    return "undetermined"


def histogram(reviews_by_id, review_ids, assignments, field):
    buckets = {}
    for rid in sorted(review_ids):
        value = reviews_by_id[rid]["metadata"].get(field)
        # Multi-valued metadata is represented as one explicit value, not double-counted.
        key = json.dumps(value, sort_keys=True, ensure_ascii=False)
        if key not in buckets:
            buckets[key] = {"value": value, "review_count": 0,
                            "sentiment_counts": {s: 0 for s in SENTIMENTS}}
        buckets[key]["review_count"] += 1
        buckets[key]["sentiment_counts"][assignments[rid]["sentiment"]] += 1
    return sorted(buckets.values(), key=lambda b: (-b["review_count"], json.dumps(b["value"], ensure_ascii=False)))


def build_report(reviews, nodes, occurrences, vectors, vocabulary, groups, rules, args, source_rows, sheet_name, elapsed):
    reviews_by_id = {r["id"]: r for r in reviews}
    for occurrence in occurrences:
        occurrence["sentiment"] = rules.analyze(occurrence["text"])
    topics = []
    for group in groups:
        occurrence_ids = [i for node_index in group for i in nodes[node_index]["occurrences"]]
        if not occurrence_ids:
            continue
        centroid = defaultdict(float)
        for ni in group:
            for term, weight in vectors[ni].items():
                centroid[term] += weight / len(group)
        norm = math.sqrt(sum(v * v for v in centroid.values()))
        ranked_terms = sorted(centroid.items(), key=lambda item: (-item[1], vocabulary[item[0]]))
        keywords, title_terms = [], []
        for term, weight in ranked_terms[:12]:
            keywords.append({"term": vocabulary[term], "weight": round(weight, 6)})
        for term, _ in ranked_terms:
            phrase = vocabulary[term]
            words = set(phrase.split())
            if any(words <= set(t.split()) or set(t.split()) <= words for t in title_terms):
                continue
            title_terms.append(phrase)
            if len(title_terms) == 3:
                break
        scores = {ni: sum(value * centroid[term] for term, value in vectors[ni].items()) / norm if norm else 0.0 for ni in group}
        for oi in occurrence_ids:
            occurrences[oi]["cosine_to_topic_centroid"] = round(scores[occurrences[oi]["node_index"]], 6)
        assignments = {}
        for oi in occurrence_ids:
            occurrence = occurrences[oi]
            rid = occurrence["review_id"]
            assignment = assignments.setdefault(rid, {"review_id": rid, "fragment_ids": [], "_labels": [], "request": False})
            assignment["fragment_ids"].append(occurrence["fragment_id"])
            assignment["_labels"].append(occurrence["sentiment"]["label"])
            assignment["request"] |= occurrence["sentiment"]["request"]
        for assignment in assignments.values():
            assignment["sentiment"] = combined_sentiment(assignment.pop("_labels"))
        counts = Counter(a["sentiment"] for a in assignments.values())
        ranked_nodes = sorted(group, key=lambda ni: (-scores[ni], ni))
        examples = []
        for ni in ranked_nodes[:5]:
            oi = nodes[ni]["occurrences"][0]
            occurrence = occurrences[oi]
            examples.append({"review_id": occurrence["review_id"], "fragment_id": occurrence["fragment_id"],
                             "text": occurrence["text"], "sentiment": occurrence["sentiment"]["label"],
                             "centroid_similarity": round(scores[ni], 6)})
        topics.append({
            "title": " / ".join(title_terms), "title_method": "top_tfidf_terms_no_generated_text",
            "keywords": keywords, "review_count": len(assignments),
            "review_share_of_all": round(len(assignments) / len(reviews), 6) if reviews else 0.0,
            "fragment_count": len(occurrence_ids), "unique_fragment_count": len(group),
            "small_topic": len(assignments) < args.small_topic_reviews,
            "sentiment_counts": {s: counts[s] for s in SENTIMENTS},
            "negative_share_of_topic": round(counts["negative"] / len(assignments), 6),
            "reviews_with_any_negative": counts["negative"] + counts["mixed"],
            "request_review_count": sum(a["request"] for a in assignments.values()),
            "breakdowns": {field: histogram(reviews_by_id, assignments.keys(), assignments, field) for field in args.group_by},
            "examples": examples, "reviews": [assignments[rid] for rid in sorted(assignments)],
            "_nodes": group, "_occurrences": occurrence_ids,
        })
    topics.sort(key=lambda t: (-t["review_count"], -t["unique_fragment_count"], t["title"], t["_nodes"][0]))
    assigned_review_ids = set()
    for i, topic in enumerate(topics, 1):
        topic["topic_id"] = f"T{i:04d}"
        for oi in topic.pop("_occurrences"):
            occurrences[oi]["topic_id"] = topic["topic_id"]
        topic.pop("_nodes")
        for assignment in topic["reviews"]:
            assigned_review_ids.add(assignment["review_id"])
    unassigned_review_ids = set()
    for occurrence in occurrences:
        node = occurrence.pop("node_index")
        if "topic_id" not in occurrence:
            occurrence["topic_id"] = None
            occurrence["unassigned_reason"] = "no_topic_tokens" if not nodes[node]["tokens"] else "no_terms_survived_tfidf_filters"
            unassigned_review_ids.add(occurrence["review_id"])
    fragment_ids_by_review = defaultdict(list)
    for occurrence in occurrences:
        fragment_ids_by_review[occurrence["review_id"]].append(occurrence["fragment_id"])
    for review in reviews:
        review["fragment_ids"] = fragment_ids_by_review[review["id"]]
        if review["id"] not in assigned_review_ids:
            unassigned_review_ids.add(review["id"])
    warnings = [
        "Clusters reflect lexical similarity, not guaranteed semantic topics; human review has not been performed.",
        "Sentiment is a limited English lexicon/rule baseline, not a validated accuracy estimate. No evidence means undetermined, not neutral.",
        "Clause splitting may leave multiple aspects in one fragment or separate a pronoun from its subject.",
        "A fragment receives one cluster. A review can receive several topics through different fragments.",
        "Rare vocabulary removed by min_df and generic statements remain visible as unassigned fragments.",
        "Topic IDs and boundaries may change when the corpus or threshold changes.",
    ]
    if len(topics) > max(30, len(reviews) // 5):
        warnings.append("Many small clusters were produced. Inspect examples and compare higher distance thresholds; do not discard rare complaints automatically.")
    metadata_conflicts = [r["id"] for r in reviews if any(isinstance(r["metadata"].get(f), list) for f in args.group_by)]
    missing_breakdowns = [f for f in args.group_by if not any(f in r["metadata"] for r in reviews)]
    if missing_breakdowns:
        warnings.append("Breakdown columns not present in source: " + ", ".join(missing_breakdowns) + ". Their values are reported as null.")
    if metadata_conflicts:
        warnings.append(f"{len(metadata_conflicts)} review IDs have multi-valued breakdown metadata; the full value lists are preserved.")
    return {
        "schema_version": "1.0", "program_version": VERSION,
        "source": {"filename": Path(args.input).name, "sha256": hashlib.sha256(Path(args.input).read_bytes()).hexdigest(),
                   "sheet": sheet_name, "data_row_count": source_rows, "unique_review_count": len(reviews)},
        "method": {
            "topic_algorithm": "TF-IDF + cosine distance + agglomerative average linkage (nearest-neighbor chain)",
            "neural_networks": False, "external_api_calls": False, "pretrained_models": False,
            "input_columns_for_clustering": [args.text_column], "language_rules": "English",
            "tf": "1 + ln(count)", "idf": "1 + ln((1 + unique_fragment_count)/(1 + document_frequency))",
            "normalization": "L2", "ngrams": [1, 2],
            "ngram_note": "Bigrams join adjacent retained normalized tokens; stopwords have already been removed.",
            "clustering_unit": "unique whitespace-normalized lowercased fragment; all original occurrences retained",
            "distance_threshold": args.threshold, "min_df": args.min_df, "max_df_ratio": args.max_df,
            "max_features": args.max_features, "vocabulary_size": len(vocabulary),
            "topic_stopwords": sorted(STOPWORDS), "topic_word_aliases": ALIASES,
            "max_fragment_words": args.max_fragment_words, "small_topic_review_threshold": args.small_topic_reviews,
            "sentiment_method": "editable phrase/word rules with local negation handling",
            "sentiment_rules_sha256": hashlib.sha256(json.dumps(rules.configuration, sort_keys=True).encode()).hexdigest(),
            "sentiment_rules": rules.configuration,
            "neutral_policy": "Reserved for future explicit neutral annotations; unmatched text is undetermined.",
            "review_topic_sentiment_policy": "Any positive and negative evidence gives mixed; otherwise known polarity; otherwise undetermined.",
            "request_policy": "Independent flag; a request is not automatically negative.",
            "counting_unit": "unique (review_id, topic_id)",
            "similarity_note": "Cosine similarity is lexical similarity, not a confidence or probability estimate.",
            "reproducibility": "No random sampling. Same input, parameters, rules and runtime give the same assignments.",
        },
        "summary": {
            "topic_count": len(topics), "fragment_count": len(occurrences), "unique_fragment_count": len(nodes),
            "assigned_review_count": len(assigned_review_ids),
            "fully_unassigned_review_count": len(set(reviews_by_id) - assigned_review_ids),
            "reviews_with_unassigned_content": len(unassigned_review_ids),
            "unassigned_fragment_count": sum(o["topic_id"] is None for o in occurrences),
            "review_topic_assignment_count": sum(t["review_count"] for t in topics),
            "small_topic_count": sum(t["small_topic"] for t in topics),
            "clustering_seconds": round(elapsed, 3),
        },
        "warnings": warnings, "topics": topics, "reviews": reviews, "fragments": occurrences,
    }


def validate_report(report):
    reviews = {r["id"]: r for r in report["reviews"]}
    fragments = {f["fragment_id"]: f for f in report["fragments"]}
    seen_pairs = set()
    for topic in report["topics"]:
        if sum(topic["sentiment_counts"].values()) != topic["review_count"]:
            raise AssertionError("Sentiment totals do not reconcile.")
        if len(topic["reviews"]) != topic["review_count"]:
            raise AssertionError("Review count does not reconcile.")
        for assignment in topic["reviews"]:
            pair = assignment["review_id"], topic["topic_id"]
            if pair in seen_pairs:
                raise AssertionError("Duplicate review/topic pair.")
            seen_pairs.add(pair)
            for fid in assignment["fragment_ids"]:
                f = fragments[fid]
                if f["review_id"] != pair[0] or f["topic_id"] != pair[1]:
                    raise AssertionError("Broken fragment reference.")
        for buckets in topic["breakdowns"].values():
            if sum(b["review_count"] for b in buckets) != topic["review_count"]:
                raise AssertionError("Metadata breakdown does not reconcile.")
    for f in fragments.values():
        if reviews[f["review_id"]]["text"][f["start"]:f["end"]] != f["text"]:
            raise AssertionError("Fragment is not an exact source substring.")
    for r in reviews.values():
        if any(fid not in fragments or fragments[fid]["review_id"] != r["id"] for fid in r["fragment_ids"]):
            raise AssertionError("Broken review reference.")


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("input", help="Input .xlsx workbook")
    parser.add_argument("-o", "--output", default="topics.json", help="Output JSON file")
    parser.add_argument("--sheet", default=None, help="Worksheet name; default: first worksheet")
    parser.add_argument("--id-column", default="ID")
    parser.add_argument("--text-column", default="Customer Feedback")
    parser.add_argument("--group-by", nargs="*", default=["Country", "Engine Type"])
    parser.add_argument("--threshold", type=float, default=0.94, help="Average cosine distance cut in [0,1); higher merges more topics")
    parser.add_argument("--min-df", type=int, default=2, help="Minimum unique-fragment document frequency; use 1 to retain singleton vocabulary")
    parser.add_argument("--max-df", type=float, default=0.75)
    parser.add_argument("--max-features", type=int, default=10000)
    parser.add_argument("--max-fragment-words", type=int, default=65)
    parser.add_argument("--small-topic-reviews", type=int, default=3, help="Flag small topics; never delete them")
    parser.add_argument("--max-memory-mb", type=int, default=768, help="Approximate similarity allocation budget, not a total process memory guarantee")
    parser.add_argument("--rules", default=str(Path(__file__).with_name("sentiment_rules_en.json")))
    args = parser.parse_args(argv)
    if not 0 <= args.threshold < 1:
        parser.error("--threshold must be in [0, 1).")
    if not 0 < args.max_df <= 1:
        parser.error("--max-df must be in (0, 1].")
    for name in ("min_df", "max_features", "max_fragment_words", "small_topic_reviews", "max_memory_mb"):
        if getattr(args, name) < 1:
            parser.error(f"--{name.replace('_', '-')} must be positive.")
    if Path(args.output).resolve() == Path(args.input).resolve():
        parser.error("Output must not overwrite the source workbook.")
    if Path(args.output).resolve() == Path(args.rules).resolve():
        parser.error("Output must not overwrite the sentiment rules.")
    return args


def run(args):
    started = time.perf_counter()
    reviews, sheet_name, source_rows = read_reviews(args.input, args.sheet, args.id_column, args.text_column)
    log(f"Read {source_rows} rows, {len(reviews)} unique review IDs.")
    nodes, occurrences = split_fragments(reviews, args.max_fragment_words)
    log(f"Extracted {len(occurrences)} fragments ({len(nodes)} unique).")
    vectors, vocabulary = tfidf(nodes, args.min_df, args.max_df, args.max_features)
    valid_indices = [i for i, v in enumerate(vectors) if v]
    log(f"TF-IDF: {len(vocabulary)} terms; clustering {len(valid_indices)} nonempty unique fragments.")
    matrix = cosine_matrix([vectors[i] for i in valid_indices], args.max_memory_mb)
    groups = [[valid_indices[i] for i in group] for group in average_linkage(matrix, args.threshold)]
    del matrix
    elapsed = time.perf_counter() - started
    log(f"Created {len(groups)} clusters at distance threshold {args.threshold}.")
    configuration = json.loads(Path(args.rules).read_text(encoding="utf-8"))
    report = build_report(reviews, nodes, occurrences, vectors, vocabulary, groups,
                          SentimentRules(configuration), args, source_rows, sheet_name, elapsed)
    validate_report(report)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.name + ".tmp")
    temporary.write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    temporary.replace(output)
    log(f"Saved {output.resolve()} ({output.stat().st_size / 1024**2:.1f} MB).")
    log(json.dumps(report["summary"], ensure_ascii=False))
    return report


def main(argv=None):
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(errors="backslashreplace")
    try:
        run(parse_args(argv))
        return 0
    except (ValueError, KeyError, FileNotFoundError, OSError, MemoryError, BadZipFile, InvalidFileException) as error:
        log(f"ERROR: {error}")
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
